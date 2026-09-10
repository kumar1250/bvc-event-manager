import csv
import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet


def _build_rows(submissions):
    """Returns (columns, rows) with dynamic columns derived from the submitted forms' fields."""
    base_columns = ["Registration ID", "Event", "Status", "Submitted At"]
    dynamic_labels = []
    seen = set()
    for s in submissions:
        for label in s.as_dict().keys():
            if label not in seen:
                seen.add(label)
                dynamic_labels.append(label)

    columns = base_columns + dynamic_labels
    rows = []
    for s in submissions:
        answers = s.as_dict()
        row = [
            s.registration_id, s.event.name, s.get_status_display(),
            s.submitted_at.strftime("%Y-%m-%d %H:%M"),
        ] + [answers.get(label, "") for label in dynamic_labels]
        rows.append(row)
    return columns, rows


def export_csv(submissions, filename="registrations.csv"):
    columns, rows = _build_rows(submissions)
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow(columns)
    writer.writerows(rows)
    return response


def export_excel(submissions, filename="registrations.xlsx"):
    columns, rows = _build_rows(submissions)
    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    for row in rows:
        ws.append(row)
    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_pdf(submissions, filename="registrations.pdf", title="Registrations"):
    columns, rows = _build_rows(submissions)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()

    elements = [Paragraph(title, styles["Title"])]
    table_data = [columns] + [[str(v) for v in row] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
